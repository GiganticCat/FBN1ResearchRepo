#!/usr/bin/env python3
"""
19_supplementary_tables.py — PHASE 6. Manuscript supplementary tables as one XLSX workbook.

Sheets:
  README            what each sheet is, where it came from, and the caveats a reader needs first
  S1_variants       the 4,451-variant master table: labels, domain/site annotation, scores,
                    structural metrics and FoldX ddG, joined into one row per variant
  S2_structural     the 751 structurally covered variants only, with their measured geometry
  S3_cmm_sites      the 16 CheckMyMetal calcium sites (experimental controls + AF3 models)
  S4_statistics     every statistical test, with effect size, CI and BH-corrected q
  S5_calibration    AF3 model vs experimental reference, per cbEGF domain
  S6_transplant     leave-one-out homology transplant of Ca2+, donor -> target
  S7_dictionary     plain-English definition of every column in S1

The workbook is a view of `data/processed/`, never a second copy of the analysis: every value
is read from those files, nothing is recomputed here. That is deliberate — a supplementary
table that recalculates its own numbers is a second chance to disagree with the paper.

Writes: results/supplementary_tables.xlsx, logs/19_supplementary_*.log
"""

from __future__ import annotations

import re
import sys
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import figstyle as fs  # noqa: E402

OUT = fs.RESULTS / "supplementary_tables.xlsx"

S1_COLUMNS = [
    "variation_id", "hgvs_c", "hgvs_p", "position", "wt_aa", "mut_aa",
    "clinvar_class", "review_status", "stars", "set_primary", "set_sensitivity",
    "gnomad_joint_af", "faf95_grpmax", "am_pathogenicity", "am_class",
    "domain", "domain_kind", "cbegf_index", "in_cbegf", "in_neonatal_region",
    "site_class", "site_label", "ca_ligand_role", "is_ca_consensus", "is_ca_sidechain_ligand",
    "is_cysteine_site", "cys_role", "disulfide_partner", "cys_removing", "cys_creating",
    "in_cbegf_cys", "vcep_pm1_category", "vcep_pm1_strength", "is_critical_residue",
    "structural_coverage", "template_pdb", "rsa_pct", "buried", "secondary_structure",
    "dist_to_nearest_ca_A", "is_direct_ca_ligand", "ss_bond_dist_A",
    "foldx_status", "ddG_kcal_mol", "ddG_disulfide", "ddG_vdw_clash",
    "mechanism_class",
]

S2_COLUMNS = [
    "variation_id", "position", "wt_aa", "mut_aa", "set_primary", "cbegf_index",
    "mechanism_class", "template_pdb", "residue", "sasa_abs_A2", "rsa_pct", "buried",
    "secondary_structure", "dist_to_nearest_ca_A", "is_direct_ca_ligand",
    "disulfide_partner", "ss_bond_dist_A", "ddG_kcal_mol", "ddG_disulfide",
    "ddG_vdw_clash", "ddG_electrostatics", "ddG_solvation_polar",
    "ddG_backbone_hbond", "ddG_sidechain_hbond", "am_pathogenicity",
]

# The Phase 2 data dictionary documents only the columns `05_normalize_variants.py` produced.
# S1 also carries the Phase 3 annotation and the Phase 5 structural / FoldX columns, so their
# definitions live here. Each was read off the script that writes the column, not inferred from
# its name. `main()` refuses to build the workbook if any S1 column is still undocumented —
# a supplementary table with unexplained columns is not a supplementary table.
EXTRA_DEFINITIONS = [
    ("Labelled sets (Phase 2)", "set_primary", "str",
     "Label under the primary tier (ClinVar review status >=2 stars): pathogenic, benign, vus, "
     "or empty when the variant does not meet any set's criteria. Sets are disjoint."),
    ("Labelled sets (Phase 2)", "set_sensitivity", "str",
     "Same, under the relaxed >=1-star tier. Used only for sensitivity analysis."),

    ("Domain annotation (Phase 3)", "domain", "str",
     "UniProt feature description containing this position, e.g. 'EGF-like 13; calcium-binding'. "
     "Empty outside any annotated domain."),
    ("Domain annotation (Phase 3)", "domain_kind", "str",
     "cbEGF | EGF | TB | outside_domain. TB includes the hybrid domains, which UniProt types "
     "as TB features."),
    ("Domain annotation (Phase 3)", "cbegf_index", "int",
     "1-43 for the calcium-binding EGF domains, in sequence order. Empty elsewhere."),
    ("Domain annotation (Phase 3)", "in_cbegf", "bool",
     "Position lies inside a cbEGF domain. NOTE: domain-resident is not the same as "
     "calcium-coordinating — see is_ca_consensus and is_direct_ca_ligand."),
    ("Domain annotation (Phase 3)", "in_neonatal_region", "bool",
     "Position is in residues 659-1362, the exon 24-32 neonatal / severe cluster."),

    ("Calcium site (Phase 3)", "site_class", "str",
     "ca_consensus (a consensus position of the calcium motif) | ca_backbone (a position whose "
     "backbone carbonyl coordinates the ion) | empty."),
    ("Calcium site (Phase 3)", "site_label", "str",
     "Which consensus position: ca_D1, ca_DN2, ca_EH, ca_DN_bOH (the beta-hydroxylation site), "
     "ca_YF (aromatic core), or ca_backbone_1..3."),
    ("Calcium site (Phase 3)", "ca_ligand_role", "str",
     "sidechain_ligand | backbone_ligand | consensus_only | aromatic_core. Derived from the "
     "motif, independent of any structure."),
    ("Calcium site (Phase 3)", "is_ca_consensus", "bool",
     "site_class == ca_consensus. Annotation-level; available for all 4,451 variants."),
    ("Calcium site (Phase 3)", "is_ca_sidechain_ligand", "bool",
     "ca_ligand_role == sidechain_ligand."),

    ("Cysteines and disulfides (Phase 3)", "is_cysteine_site", "bool",
     "Wild-type residue is cysteine."),
    ("Cysteines and disulfides (Phase 3)", "cys_role", "str",
     "C1-C6, the cysteine's index within its cbEGF domain. Empty for non-cbEGF cysteines."),
    ("Cysteines and disulfides (Phase 3)", "disulfide_partner", "int",
     "P35555 position of the partner cysteine. cbEGF domains pair C1-C3, C2-C4, C5-C6 without "
     "exception across all 129 cbEGF disulfides."),
    ("Cysteines and disulfides (Phase 3)", "cys_removing", "bool",
     "Wild-type is Cys and the substitution is not — a disulfide is lost. The hallmark Marfan "
     "missense mechanism."),
    ("Cysteines and disulfides (Phase 3)", "cys_creating", "bool",
     "Wild-type is not Cys and the substitution is — an unpaired thiol is introduced."),
    ("Cysteines and disulfides (Phase 3)", "in_cbegf_cys", "bool",
     "Position is one of the 258 conserved cbEGF cysteines."),

    ("ClinGen VCEP PM1 (Phase 3)", "vcep_pm1_category", "str",
     "Which PM1 category the position falls in under the ClinGen FBN1 VCEP rules, e.g. "
     "cbEGF_cysteine_removing, calcium_consensus_residue, interdomain_packing_glycine."),
    ("ClinGen VCEP PM1 (Phase 3)", "vcep_pm1_strength", "str",
     "PM1_Strong | PM1 | not_applied | empty. not_applied marks the published tolerated "
     "Asn>Ser exception at the second D/N position."),
    ("ClinGen VCEP PM1 (Phase 3)", "is_critical_residue", "bool",
     "PM1 applies at some strength. Used only in the enrichment analysis, which is partly "
     "circular because ClinVar's labels were themselves assigned using PM1."),

    ("Structural metrics (Phase 5)", "structural_coverage", "bool",
     "The position falls inside one of the four calcium-bound constructs and carries measured "
     "geometry. True for 751 of 4,451 variants."),
    ("Structural metrics (Phase 5)", "template_pdb", "str",
     "Structure the measurements were taken from: 2W86, 1UZJ, 1LMJ or 1EMN. Renumbered into "
     "P35555 numbering via SIFTS before measurement."),
    ("Structural metrics (Phase 5)", "residue", "str",
     "Residue found at that position in the structure. Verified against P35555 before any "
     "measurement was recorded."),
    ("Structural metrics (Phase 5)", "sasa_abs_A2", "float",
     "Absolute solvent-accessible surface area of the wild-type residue, A^2 (freesasa)."),
    ("Structural metrics (Phase 5)", "rsa_pct", "float",
     "Relative solvent accessibility, % of the residue's maximum accessible area."),
    ("Structural metrics (Phase 5)", "buried", "bool",
     "rsa_pct < 20%."),
    ("Structural metrics (Phase 5)", "secondary_structure", "str",
     "DSSP class at that position."),
    ("Structural metrics (Phase 5)", "dist_to_nearest_ca_A", "float",
     "Distance from the nearest side-chain or backbone atom to the nearest Ca2+ ion, A. In AF3 "
     "models the ion was PLACED by cofolding, not observed."),
    ("Structural metrics (Phase 5)", "is_direct_ca_ligand", "bool",
     "The residue has an oxygen or nitrogen within 3.2 A of a Ca2+ ion — i.e. it actually "
     "coordinates the metal in the structure. Stricter and structure-based, unlike "
     "is_ca_consensus."),
    ("Structural metrics (Phase 5)", "ss_bond_dist_A", "float",
     "S-S distance to the partner cysteine, A. Only for cysteine positions with coverage."),

    ("FoldX (Phase 5)", "foldx_status", "str",
     "ok | no_structure | parse_failed | MISSING. Every covered variant returned ok; nothing "
     "was silently dropped."),
    ("FoldX (Phase 5)", "ddG_kcal_mol", "float",
     "Predicted change in folding free energy, kcal/mol. Positive = destabilising. FoldX 5.1, "
     "RepairPDB then BuildModel; mutation strings written in author numbering via SIFTS."),
    ("FoldX (Phase 5)", "ddG_disulfide", "float",
     "The disulfide term of the FoldX decomposition. Isolates staple loss from general strain."),
    ("FoldX (Phase 5)", "ddG_vdw_clash", "float",
     "The van der Waals clash term of the FoldX decomposition."),
    ("FoldX (Phase 5)", "ddG_electrostatics", "float", "Electrostatics term."),
    ("FoldX (Phase 5)", "ddG_solvation_polar", "float", "Polar solvation term."),
    ("FoldX (Phase 5)", "ddG_backbone_hbond", "float", "Backbone hydrogen-bond term."),
    ("FoldX (Phase 5)", "ddG_sidechain_hbond", "float", "Side-chain hydrogen-bond term."),

    ("Derived for the figures (Phase 6)", "mechanism_class", "str",
     "ca_ligand | cys_removing | other_cbegf | not_in_cbEGF. The grouping used in Figures 2-4, "
     "defined exactly as in 12_statistics.py: ca_ligand is the structure-measured "
     "is_direct_ca_ligand, NOT the annotation-level is_ca_consensus."),
]

README_ROWS = [
    ("FBN1 / Marfan structural bioinformatics — supplementary tables", ""),
    ("", ""),
    ("Canonical protein", "UniProt P35555 (fibrillin-1), 2,871 aa"),
    ("Canonical transcript", "NM_000138.5 (MANE Select)"),
    ("ClinVar release", "variant_summary 2026-07-28"),
    ("UniProt entry version", "264"),
    ("gnomAD dataset", "gnomad_r4 (GRCh38)"),
    ("AlphaMissense", "Zenodo 8208688 v1.0.0 (Cheng et al. 2023, CC BY-NC-SA 4.0)"),
    ("", ""),
    ("READ FIRST", "results/LIMITATIONS.md. Three caveats govern every number below:"),
    ("1. Enrichment is partly circular",
     "ClinVar's pathogenic calls used ClinGen PM1 — the same critical-residue rule the "
     "enrichment tests. The AlphaMissense contrast is the orthogonal check."),
    ("2. The benign set is small and cannot be enlarged",
     "34 benign variants at >=2 stars. Of 2,783 gnomAD missense variants only 24 clear the "
     "benign filtering-AF threshold and all 24 were already in the ClinVar set."),
    ("3. Calcium in predicted models was placed, not observed",
     "AlphaFold models contain no metals. Every ion in an AF3 model was cofolded; median "
     "displacement vs X-ray references 0.32 A (S5). Homology transplant is ~4x worse (S6)."),
    ("", ""),
    ("Structural coverage", "751 of 4,451 variants (17%) fall inside the four calcium-bound "
                            "constructs and carry measured geometry. The other 3,700 carry "
                            "annotation and AlphaMissense scores only, and are marked "
                            "structural_coverage = False rather than left blank."),
    ("Numbering", "Every position is P35555 numbering. Structure files were renumbered from "
                  "PDB author numbering via SIFTS in Phase 4 before any measurement."),
    ("", ""),
    ("SHEETS", ""),
    ("S1_variants", "Master table, one row per ClinVar GRCh38 missense variant (4,451)."),
    ("S2_structural", "The 751 structurally covered variants with their measured geometry."),
    ("S3_cmm_sites", "16 CheckMyMetal calcium sites: 8 experimental controls, 8 AF3 models."),
    ("S4_statistics", "All statistical tests with effect size, 95% CI and BH-corrected q."),
    ("S5_calibration", "AF3 model vs experimental reference, per cbEGF domain (40 rows)."),
    ("S6_transplant", "Leave-one-out Ca2+ transplant, donor -> target (56 rows)."),
    ("S7_dictionary", "Column definitions for S1."),
]


def parse_dictionary(md: Path) -> pd.DataFrame:
    """Lift the pipe-table rows out of the Phase 2 data dictionary into a flat sheet."""
    rows, section = [], ""
    for line in md.read_text(encoding="utf-8").splitlines():
        if line.startswith("## "):
            section = line[3:].strip()
        m = re.match(r"^\|\s*`([^`]+)`\s*\|\s*([^|]*?)\s*\|\s*(.*?)\s*\|\s*$", line)
        if m:
            rows.append({"section": section, "column": m.group(1),
                         "type": m.group(2), "meaning": m.group(3)})
    return pd.DataFrame(rows)


def autoformat(writer, sheet: str, df: pd.DataFrame, wrap_cols: set[str] = frozenset()) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = writer.sheets[sheet]
    head_fill = PatternFill("solid", fgColor="EFEEE9")
    for j, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=j)
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        body = df[col].astype(str)
        width = max(len(str(col)), int(body.str.len().quantile(0.95)) if len(body) else 0) + 2
        if col in wrap_cols:
            width = 70
            for i in range(2, len(df) + 2):
                ws.cell(row=i, column=j).alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = min(max(width, 9), 70)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main() -> int:
    log = [f"Phase 6 — supplementary tables — {datetime.now(timezone.utc).isoformat()}"]

    df = fs.load_merged()
    df["mechanism_class"] = fs.mech_group(df).fillna("not_in_cbEGF")

    missing = [c for c in S1_COLUMNS if c not in df.columns]
    if missing:
        raise SystemExit(f"S1 columns absent from the processed tables: {missing}")

    s1 = df[S1_COLUMNS].sort_values("position").reset_index(drop=True)
    s2 = (df[df.structural_coverage][S2_COLUMNS]
          .sort_values(["template_pdb", "position"]).reset_index(drop=True))
    s3 = pd.read_csv(fs.PROC / "cmm_sites.tsv", sep="\t")
    s4 = pd.read_csv(fs.RESULTS / "phase5_statistics.tsv", sep="\t")
    s5 = pd.read_csv(fs.RESULTS / "phase4_calibration.tsv", sep="\t")
    s6 = pd.read_csv(fs.RESULTS / "phase4_transplant_loo.tsv", sep="\t")
    s7 = pd.concat([parse_dictionary(fs.PROC / "variant_master_dictionary.md"),
                    pd.DataFrame(EXTRA_DEFINITIONS,
                                 columns=["section", "column", "type", "meaning"])],
                   ignore_index=True)
    s7 = s7[s7.column.isin(set(S1_COLUMNS) | set(S2_COLUMNS))].drop_duplicates("column")
    undocumented = [c for c in S1_COLUMNS + S2_COLUMNS if c not in set(s7.column)]
    if undocumented:
        raise SystemExit(f"columns in the supplementary tables with no definition: "
                         f"{undocumented}")
    s7 = s7.reset_index(drop=True)

    if len(s1) != 4451:
        raise SystemExit(f"S1 has {len(s1)} rows, expected the full 4,451 variant table")
    if len(s2) != int(df.structural_coverage.sum()):
        raise SystemExit("S2 row count does not match the structurally covered set")

    readme = pd.DataFrame(README_ROWS, columns=["Item", "Detail"])

    with pd.ExcelWriter(OUT, engine="openpyxl") as w:
        sheets = [
            ("README", readme, {"Detail"}),
            ("S1_variants", s1, set()),
            ("S2_structural", s2, set()),
            ("S3_cmm_sites", s3, set()),
            ("S4_statistics", s4, {"note", "test"}),
            ("S5_calibration", s5, set()),
            ("S6_transplant", s6, set()),
            ("S7_dictionary", s7, {"meaning"}),
        ]
        for name, frame, wrap in sheets:
            frame.to_excel(w, sheet_name=name, index=False)
            autoformat(w, name, frame, wrap)
            log.append(f"  {name:16} {len(frame):5} rows x {len(frame.columns):2} cols")
            print(f"  {name:16} {len(frame):5} rows × {len(frame.columns):2} cols")

    size_kb = OUT.stat().st_size / 1024
    log.append(f"wrote {OUT.relative_to(fs.ROOT)} ({size_kb:.0f} kB)")
    print(f"  wrote {OUT.relative_to(fs.ROOT)}  ({size_kb:.0f} kB)")
    print(f"  log -> {fs.write_log('19_supplementary', log).name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
