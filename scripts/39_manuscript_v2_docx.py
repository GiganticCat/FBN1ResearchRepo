#!/usr/bin/env python3
"""
39_manuscript_v2_docx.py — render the v2 manuscript draft to a styled Word document.

`results/manuscript_v2_draft.md` stays the source of truth: edit the markdown, re-run this, and
the document follows. Nothing is retyped here, so the two cannot drift.

What this adds over the markdown:
  * the six figures placed inline, each after the section that discusses it, rather than as a
    list of legends at the end -- the draft is for reading, not for submission formatting
  * the legend text pulled from the markdown's own legend section and used as the caption, so
    there is exactly one copy of each legend
  * document styling: a serif body at manuscript size, coloured headings, captioned figures

Writes: results/FBN1_manuscript_v2_draft.docx
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "results" / "manuscript_v2_draft.md"
FIGDIR = ROOT / "figures"
OUT = ROOT / "results" / "FBN1_manuscript_v2_draft.docx"

# Which figure follows which results subsection, and which file it is. The manuscript numbers
# figures 1-6 in reading order; the filenames carry their own history and do not match, which is
# exactly why this mapping is written down rather than inferred from a sort.
PLACEMENT = [
    ("A structural basis for most of the variant set", 1, "v2_fig3_landscape"),
    ("Calcium-site variants do not destabilise the domain", 2, "v2_fig1_stability"),
    ("The result is not an artefact of where these residues sit", 3, "v2_fig4_burial"),
    ("The damage appears in the ion's coordination", 4, "v2_fig2_geometry"),
    ("The modelled calcium sites reproduce experimental ones", 5, "v2_fig5_validation"),
    ("Enrichment is consistent but partly circular", 6, "v2_fig6_enrichment"),
]

INK = (0x1A, 0x1F, 0x21)
ACCENT = (0x0F, 0x5C, 0x44)       # the calcium green, darkened for text contrast on white
MUTED = (0x5A, 0x63, 0x5F)
BODY_FACE = "Cambria"             # serif, ships with Word; Georgia is the fallback
SANS_FACE = "Calibri"


def parse_legends(text: str) -> dict[int, str]:
    """Pull '**Fig. N | Title.** body' (or '**Figure N |**') blocks out of the legends section."""
    out: dict[int, str] = {}
    sec = text.split("## Figure legends", 1)
    if len(sec) < 2:
        return out
    body = sec[1].split("\n---", 1)[0]
    for block in re.split(r"\n\s*\n", body):
        block = block.strip()
        m = re.match(r"\*\*(?:Figure|Fig\.) (\d+) \|", block)
        if m:
            out[int(m.group(1))] = " ".join(block.split())
    return out


def add_runs(par, text: str, size, base_color, italic_all: bool = False) -> None:
    """Emit **bold**, *italic* and `code` as separate runs."""
    from docx.shared import Pt, RGBColor
    for piece in re.split(r"(\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*)|`[^`]+`)", text):
        if not piece:
            continue
        bold = italic = code = False
        if piece.startswith("**") and piece.endswith("**"):
            piece, bold = piece[2:-2], True
        elif piece.startswith("`") and piece.endswith("`"):
            piece, code = piece[1:-1], True
        elif piece.startswith("*") and piece.endswith("*"):
            piece, italic = piece[1:-1], True
        run = par.add_run(piece)
        run.bold = bold
        run.italic = italic or italic_all
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(*base_color)
        run.font.name = "Consolas" if code else BODY_FACE
        if code:
            run.font.size = Pt(size - 1)
    return None


def main() -> int:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:
        sys.exit("STOP: python-docx is not installed in this environment")

    if not SRC.is_file():
        sys.exit(f"STOP: {SRC.relative_to(ROOT)} not found")
    text = SRC.read_text(encoding="utf-8")
    legends = parse_legends(text)
    missing = [n for _, n, _ in PLACEMENT if n not in legends]
    if missing:
        sys.exit(f"STOP: no legend text found for figure(s) {missing}")

    doc = Document()
    sec = doc.sections[0]
    # 0.85 in rather than a reflexive 1.0: the figures are drawn at 180 mm (7.09 in) double-
    # column with 7 pt labels, so a 6.5 in text column shrinks their type to about 92% of design
    # size. 6.8 in brings that back to ~96% without the page looking unusual.
    sec.left_margin = sec.right_margin = Inches(0.85)
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    usable = sec.page_width - sec.left_margin - sec.right_margin

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FACE
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(*INK)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.18

    def para(txt="", size=11, color=INK, space_after=8, italic=False, align=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        if align is not None:
            p.alignment = align
        if txt:
            add_runs(p, txt, size, color, italic_all=italic)
        return p

    def heading(txt, level):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
        p.paragraph_format.space_after = Pt(5)
        size = 14 if level == 1 else 11.5
        colour = ACCENT if level == 1 else INK
        # Headings carry inline markup too — one of them ends in *[preliminary]*, and emitting
        # the heading as a single run printed the asterisks literally.
        for piece in re.split(r"(\*[^*]+\*)", txt):
            if not piece:
                continue
            italic = piece.startswith("*") and piece.endswith("*")
            r = p.add_run(piece[1:-1] if italic else piece)
            r.bold = not italic
            r.italic = italic
            r.font.name = SANS_FACE
            r.font.size = Pt(size - 1 if italic else size)
            r.font.color.rgb = RGBColor(*(MUTED if italic else colour))
        return p

    def figure(num: int, stem: str):
        png = FIGDIR / f"{stem}.png"
        if not png.is_file():
            sys.exit(f"STOP: missing figure image {png.relative_to(ROOT)}")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.add_run().add_picture(str(png), width=usable)
        cap = doc.add_paragraph()
        cap.paragraph_format.space_after = Pt(14)
        cap.paragraph_format.left_indent = Inches(0.12)
        cap.paragraph_format.right_indent = Inches(0.12)
        add_runs(cap, legends[num], 8.5, MUTED)

    # ------------------------------------------------------------------ walk the markdown
    lines = text.splitlines()
    i = 0
    pending_fig: tuple[int, str] | None = None
    in_legends = False
    buf: list[str] = []

    def flush():
        nonlocal buf
        if buf:
            para(" ".join(buf))
            buf = []

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()

        if line.startswith("## Figure legends"):
            # legends are emitted with their figures, so the standalone section is dropped
            in_legends = True
            i += 1
            continue
        if in_legends:
            if line.startswith("## "):
                in_legends = False
            else:
                i += 1
                continue

        if not line:
            flush()
            i += 1
            continue

        if line.startswith("# "):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(line[2:])
            r.bold = True
            r.font.name = SANS_FACE
            r.font.size = Pt(19)
            r.font.color.rgb = RGBColor(*INK)
            i += 1
            continue

        if line.startswith("### "):
            flush()
            # the previous subsection's figure goes in before the next one starts
            if pending_fig:
                figure(*pending_fig)
                pending_fig = None
            title = line[4:]
            heading(title, 2)
            for anchor, num, stem in PLACEMENT:
                if title.startswith(anchor):
                    pending_fig = (num, stem)
            i += 1
            continue

        if line.startswith("## "):
            flush()
            if pending_fig:
                figure(*pending_fig)
                pending_fig = None
            heading(line[3:], 1)
            i += 1
            continue

        if line.startswith("---"):
            flush()
            if pending_fig:
                figure(*pending_fig)
                pending_fig = None
            i += 1
            continue

        if re.match(r"^\d+\.\s", line):
            flush()
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, re.sub(r"^\d+\.\s", "", line), 11, INK)
            i += 1
            continue

        if line.startswith("- ") or line.startswith("* "):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, line[2:], 11, INK)
            i += 1
            continue

        if line.startswith("**Draft"):
            flush()
            para(line, size=9.5, color=MUTED, italic=True, space_after=12)
            i += 1
            continue

        buf.append(line)
        i += 1

    flush()
    if pending_fig:
        figure(*pending_fig)

    doc.save(OUT)
    kb = OUT.stat().st_size / 1024
    print(f"wrote {OUT.relative_to(ROOT)}  ({kb:.0f} kB)")
    print(f"  {len(PLACEMENT)} figures placed inline with captions")
    print(f"  source: {SRC.relative_to(ROOT)} — edit that and re-run")
    return 0


if __name__ == "__main__":
    sys.exit(main())
